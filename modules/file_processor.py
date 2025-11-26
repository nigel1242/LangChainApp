# modules/file_processor.py
import io
from typing import List
import os
import pandas as pd
from docx import Document as DocxDocument
from pptx import Presentation
import fitz  # PyMuPDF

TEXT_FILE_EXTS = {".txt", ".md"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
CSV_EXTS = {".csv"}
PDF_EXTS = {".pdf"}


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def extract_text_from_file(file) -> str:
    """Extract text from uploaded file-like object."""
    name = file.name.lower()
    text = ""

    if name.endswith(".pdf"):
        doc = fitz.open(stream=file.read(), filetype="pdf")
        for page in doc:
            txt = page.get_text("text")
            if txt:
                text += txt + "\n"
        doc.close()

    elif name.endswith(".docx"):
        doc = DocxDocument(file)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    elif name.endswith(".pptx"):
        prs = Presentation(file)
        slides_text = []
        for i, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                slides_text.append(f"--- Slide {i} ---\n" + "\n".join(slide_text))
        text = "\n".join(slides_text)

    elif name.endswith((".txt", ".md")):
        text = file.read().decode("utf-8", errors="ignore")

    elif name.endswith(".csv"):
        try:
            df = pd.read_csv(file)
            text = df.to_string(index=False)
        except Exception as e:
            text = f"Error reading CSV: {e}"

    return text.strip()


def extract_text_from_path(path: str) -> str:
    """Extract text from file on disk."""
    ext = _ext(path)

    if ext in TEXT_FILE_EXTS:
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            with open(path, "r", encoding="cp1252", errors="ignore") as f:
                return f.read()

    elif ext in DOCX_EXTS:
        try:
            doc = DocxDocument(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return ""

    elif ext in PPTX_EXTS:
        try:
            prs = Presentation(path)
            slides_text = []
            for i, slide in enumerate(prs.slides, start=1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    slides_text.append(f"--- Slide {i} ---\n" + "\n".join(slide_text))
            return "\n".join(slides_text)
        except Exception:
            return ""

    elif ext in CSV_EXTS:
        try:
            df = pd.read_csv(path)
            return df.to_string(index=False)
        except Exception as e:
            return f"Error reading CSV: {e}"

    elif ext in PDF_EXTS:
        try:
            doc = fitz.open(path)
            text = []
            for page in doc:
                txt = page.get_text("text")
                if txt:
                    text.append(txt)
            doc.close()
            return "\n".join(text)
        except Exception:
            return ""

    return ""
